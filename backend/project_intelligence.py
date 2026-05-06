from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any
import json
import re

try:
    from openpyxl import load_workbook  # type: ignore
except Exception:
    load_workbook = None


@dataclass
class ProjectConfig:
    key: str
    name: str
    root_path: Path


class ProjectIntelligenceService:
    def __init__(self) -> None:
        self._config = self._load_properties()
        self._default_projects = ["Cummins", "Poinier", "SenseoPod", "Golab"]

    def _load_properties(self) -> dict[str, str]:
        config_path = Path(__file__).parent / "engine.properties"
        if not config_path.exists():
            return {}
        values: dict[str, str] = {}
        for line in config_path.read_text(encoding="utf-8").splitlines():
            clean = line.strip()
            if not clean or clean.startswith("#") or "=" not in clean:
                continue
            key, value = clean.split("=", 1)
            values[key.strip()] = value.strip()
        return values

    def list_projects(self) -> list[dict[str, Any]]:
        configured: list[ProjectConfig] = []
        for key, value in self._config.items():
            if key.startswith("projects.") and key.endswith(".path"):
                project_key = key.split(".")[1]
                name = self._config.get(f"projects.{project_key}.name", project_key.title())
                configured.append(ProjectConfig(key=project_key, name=name, root_path=Path(value).expanduser()))

        if not configured:
            return [
                {
                    "id": item.lower(),
                    "name": item,
                    "path": str((Path.cwd() / "data" / item.lower()).resolve()),
                    "available": False,
                }
                for item in self._default_projects
            ]

        return [
            {
                "id": item.key,
                "name": item.name,
                "path": str(item.root_path),
                "available": item.root_path.exists(),
            }
            for item in sorted(configured, key=lambda p: p.name.lower())
        ]

    def _iter_readable_files(self, root: Path) -> list[Path]:
        if not root.exists() or not root.is_dir():
            return []
        files: list[Path] = []
        for ext in ("*.xlsx", "*.xlsm", "*.csv", "*.txt", "*.log", "*.json"):
            files.extend(root.rglob(ext))
        return sorted(files, key=lambda p: p.stat().st_mtime, reverse=True)[:100]

    def _read_excel_preview(self, path: Path) -> list[dict[str, Any]]:
        if load_workbook is None:
            return [{"sheet": "unavailable", "data": "Install openpyxl to parse Excel files."}]
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception:
            return [{"sheet": "unreadable", "data": f"Could not parse {path.name}"}]

        previews: list[dict[str, Any]] = []
        for sheet_name in workbook.sheetnames[:3]:
            sheet = workbook[sheet_name]
            rows = []
            for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                compact = [cell for cell in row if cell is not None]
                if compact:
                    rows.append(compact)
                if index >= 10:
                    break
            previews.append({"sheet": sheet_name, "data": rows})
        workbook.close()
        return previews

    def _read_text_preview(self, path: Path) -> str:
        try:
            content = path.read_text(encoding="utf-8", errors="ignore")
        except Exception:
            return ""
        return "\n".join(content.splitlines()[:20]).strip()

    def project_insights(self, project_id: str) -> dict[str, Any]:
        project = next((item for item in self.list_projects() if item["id"] == project_id), None)
        if not project:
            return {"error": "Project not found"}

        root = Path(project["path"])
        files = self._iter_readable_files(root)
        if not files:
            return {
                "project": project,
                "latest_milestone": "No files discovered",
                "latest_invoice": "No invoice data discovered",
                "cards": [],
            }

        cards: list[dict[str, Any]] = []
        milestone_hits: list[str] = []
        invoice_hits: list[str] = []

        for file in files[:20]:
            suffix = file.suffix.lower()
            if suffix in {".xlsx", ".xlsm"}:
                preview = self._read_excel_preview(file)
                text_probe = json.dumps(preview).lower()
                cards.append(
                    {
                        "title": file.name,
                        "type": "excel",
                        "updated_at": datetime.fromtimestamp(file.stat().st_mtime).isoformat(),
                        "preview": preview,
                    }
                )
            else:
                preview = self._read_text_preview(file)
                text_probe = preview.lower()
                cards.append(
                    {
                        "title": file.name,
                        "type": suffix.lstrip("."),
                        "updated_at": datetime.fromtimestamp(file.stat().st_mtime).isoformat(),
                        "preview": preview,
                    }
                )

            if "milestone" in text_probe:
                milestone_hits.append(f"{file.name}: milestone update available")
            match = re.search(r"invoice[^\n]{0,50}", text_probe)
            if match:
                invoice_hits.append(f"{file.name}: {match.group(0)[:80]}")

        return {
            "project": project,
            "latest_milestone": milestone_hits[0] if milestone_hits else "Milestone not found in sampled files",
            "latest_invoice": invoice_hits[0] if invoice_hits else "Invoice data not found in sampled files",
            "cards": cards[:10],
        }

    def _mock_qoq(self, summaries: list[dict[str, Any]]) -> list[dict[str, Any]]:
        counts = Counter(summary["card_count"] for summary in summaries)
        baseline = sum(score * count for score, count in counts.items()) or 1
        return [
            {"quarter": "Q1", "score": round(baseline * 0.82, 2)},
            {"quarter": "Q2", "score": round(baseline * 0.93, 2)},
            {"quarter": "Q3", "score": round(baseline * 1.01, 2)},
            {"quarter": "Q4", "score": round(baseline * 1.10, 2)},
        ]

    def answer_query(self, query: str, project_id: str | None = None) -> dict[str, Any]:
        if project_id:
            details = self.project_insights(project_id)
            return {
                "scope": "project",
                "project_id": project_id,
                "query": query,
                "answer": (
                    f"Project scan complete. {details.get('latest_milestone', '')}. "
                    f"{details.get('latest_invoice', '')}"
                ),
                "cards": details.get("cards", []),
            }

        project_summaries: list[dict[str, Any]] = []
        risks: list[str] = []
        for project in self.list_projects():
            if not project["available"]:
                risks.append(f"{project['name']}: folder not available")
                continue
            insight = self.project_insights(project["id"])
            project_summaries.append(
                {
                    "project": project["name"],
                    "milestone": insight.get("latest_milestone"),
                    "invoice": insight.get("latest_invoice"),
                    "card_count": len(insight.get("cards", [])),
                }
            )
            if "not found" in str(insight.get("latest_milestone", "")).lower():
                risks.append(f"{project['name']}: milestone tracking missing")

        return {
            "scope": "bu",
            "query": query,
            "answer": "BU summary generated across configured projects.",
            "summary": project_summaries,
            "qoq": self._mock_qoq(project_summaries),
            "risks": risks[:6],
        }


project_intelligence_service = ProjectIntelligenceService()
